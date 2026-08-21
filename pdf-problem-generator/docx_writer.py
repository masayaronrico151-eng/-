"""生成した問題をWord文書(.docx)に出力するモジュール"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DIFFICULTY_LABELS = {
    "easy": "基礎",
    "standard": "標準",
    "hard": "応用",
}

DIFFICULTY_COLORS = {
    "easy": RGBColor(0x22, 0x8B, 0x22),      # 緑
    "standard": RGBColor(0x1E, 0x90, 0xFF),   # 青
    "hard": RGBColor(0xDC, 0x14, 0x3C),       # 赤
}


def _set_document_style(doc: Document):
    """文書全体のスタイルを設定"""
    style = doc.styles["Normal"]
    font = style.font
    font.name = "游ゴシック"
    font.size = Pt(10.5)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)


def _add_title(doc: Document, title: str):
    """タイトルを追加"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(title)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    doc.add_paragraph()  # 空行


def _add_section_header(doc: Document, text: str):
    """セクションヘッダーを追加"""
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    para_format = para.paragraph_format
    para_format.space_after = Pt(6)


def _add_difficulty_header(doc: Document, difficulty: str):
    """難易度ヘッダーを追加"""
    label = DIFFICULTY_LABELS.get(difficulty, difficulty)
    color = DIFFICULTY_COLORS.get(difficulty, RGBColor(0, 0, 0))

    para = doc.add_paragraph()
    run = para.add_run(f"【{label}レベル】")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = color


def _add_problem(doc: Document, number: str, text: str, indent: bool = False):
    """問題を追加"""
    para = doc.add_paragraph()
    if indent:
        para.paragraph_format.left_indent = Cm(0.5)

    run = para.add_run(f"{number}  ")
    run.bold = True
    run.font.size = Pt(10.5)

    run2 = para.add_run(text)
    run2.font.size = Pt(10.5)

    para.paragraph_format.space_after = Pt(8)


def _add_answer_block(doc: Document, number: str, answer: str, explanation: str):
    """解答と解説を追加"""
    # 解答
    para = doc.add_paragraph()
    run = para.add_run(f"{number}  ")
    run.bold = True
    run.font.size = Pt(10.5)

    run2 = para.add_run(f"【解答】{answer}")
    run2.font.size = Pt(10.5)

    para.paragraph_format.space_after = Pt(2)

    # 解説
    para2 = doc.add_paragraph()
    para2.paragraph_format.left_indent = Cm(0.8)
    run3 = para2.add_run(f"【解説】{explanation}")
    run3.font.size = Pt(9.5)
    run3.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    para2.paragraph_format.space_after = Pt(12)


def write_docx(
    results: list[dict],
    output_path: str,
    subject: str,
    source_filename: str,
):
    """生成結果をWord文書に出力する

    Args:
        results: generate_all_problemsの返り値
        output_path: 出力ファイルパス
        subject: 科目名
        source_filename: 元のPDFファイル名
    """
    doc = Document()
    _set_document_style(doc)

    # === 問題ページ ===
    _add_title(doc, f"{subject} 類題集")

    # 出典情報
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = para.add_run(f"出典: {source_filename}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    doc.add_paragraph()  # 空行

    for result in results:
        original = result["original"]
        generated = result["generated"]

        # 元の問題の表示
        _add_section_header(doc, f"元の問題 {original['number']}")

        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.5)
        run = para.add_run(original["text"])
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        para.paragraph_format.space_after = Pt(12)

        # 難易度別に問題を表示
        current_difficulty = None
        problem_counter = 1

        for prob in generated:
            difficulty = prob.get("difficulty", "standard")

            if difficulty != current_difficulty:
                _add_difficulty_header(doc, difficulty)
                current_difficulty = difficulty
                problem_counter = 1

            label = DIFFICULTY_LABELS.get(difficulty, "")
            number_str = f"{label}-{problem_counter}"
            _add_problem(doc, number_str, prob["problem"], indent=True)
            problem_counter += 1

        doc.add_paragraph()  # セクション間の空白

    # === 解答ページ（改ページ） ===
    doc.add_page_break()
    _add_title(doc, f"{subject} 類題集 ― 解答・解説")

    for result in results:
        original = result["original"]
        generated = result["generated"]

        _add_section_header(doc, f"元の問題 {original['number']} の類題 解答")

        current_difficulty = None
        problem_counter = 1

        for prob in generated:
            difficulty = prob.get("difficulty", "standard")

            if difficulty != current_difficulty:
                _add_difficulty_header(doc, difficulty)
                current_difficulty = difficulty
                problem_counter = 1

            label = DIFFICULTY_LABELS.get(difficulty, "")
            number_str = f"{label}-{problem_counter}"
            _add_answer_block(
                doc,
                number_str,
                prob.get("answer", "（解答なし）"),
                prob.get("explanation", "（解説なし）"),
            )
            problem_counter += 1

        doc.add_paragraph()

    doc.save(output_path)
