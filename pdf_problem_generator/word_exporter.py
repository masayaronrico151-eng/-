"""Word文書（.docx）出力モジュール"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from .models import Difficulty, GeneratedProblem, ProblemSet


class WordExporter:
    """問題セットをWord文書として出力するクラス"""

    # 難易度ごとのカラー設定
    DIFFICULTY_COLORS = {
        Difficulty.EASY: RGBColor(0x22, 0x8B, 0x22),    # 緑
        Difficulty.NORMAL: RGBColor(0x1E, 0x90, 0xFF),   # 青
        Difficulty.HARD: RGBColor(0xDC, 0x14, 0x3C),     # 赤
    }

    def __init__(self):
        self.doc = Document()
        self._setup_styles()

    def _setup_styles(self):
        """文書の基本スタイルを設定"""
        style = self.doc.styles["Normal"]
        font = style.font
        font.name = "游明朝"
        font.size = Pt(11)

        # ページ余白の設定
        for section in self.doc.sections:
            section.top_margin = Cm(2.0)
            section.bottom_margin = Cm(2.0)
            section.left_margin = Cm(2.5)
            section.right_margin = Cm(2.5)

    def _add_title(self, title: str):
        """タイトルを追加"""
        paragraph = self.doc.add_heading(title, level=0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_section_header(self, text: str, level: int = 1):
        """セクションヘッダーを追加"""
        self.doc.add_heading(text, level=level)

    def _add_problem(self, gen: GeneratedProblem, index: int):
        """生成された1問を追加"""
        color = self.DIFFICULTY_COLORS.get(gen.difficulty, RGBColor(0, 0, 0))

        # 難易度ラベル
        p = self.doc.add_paragraph()
        run = p.add_run(f"【{gen.difficulty.label_ja}】")
        run.bold = True
        run.font.color.rgb = color
        run.font.size = Pt(12)

        # 問題番号と問題文
        p = self.doc.add_paragraph()
        run = p.add_run(f"類題 {index}. ")
        run.bold = True
        run.font.size = Pt(11)
        run = p.add_run(gen.problem_text)
        run.font.size = Pt(11)

    def _add_answer_section(self, gen: GeneratedProblem, index: int):
        """解答・解説を追加"""
        color = self.DIFFICULTY_COLORS.get(gen.difficulty, RGBColor(0, 0, 0))

        # 解答ヘッダー
        p = self.doc.add_paragraph()
        run = p.add_run(f"類題 {index}【{gen.difficulty.label_ja}】の解答")
        run.bold = True
        run.font.color.rgb = color
        run.font.size = Pt(11)

        # 解答
        p = self.doc.add_paragraph()
        run = p.add_run("【解答】")
        run.bold = True
        run = p.add_run(f"\n{gen.answer}")

        # 解説
        p = self.doc.add_paragraph()
        run = p.add_run("【解説】")
        run.bold = True
        run = p.add_run(f"\n{gen.explanation}")

        # 区切り線
        self.doc.add_paragraph("─" * 50)

    def export(
        self,
        problem_sets: list[ProblemSet],
        output_path: str | Path,
        title: str = "類題集",
        include_originals: bool = True,
    ):
        """問題セットをWord文書として出力する

        Args:
            problem_sets: 問題セットのリスト
            output_path: 出力ファイルパス
            title: 文書タイトル
            include_originals: 元問題も含めるか
        """
        output_path = Path(output_path)

        self._add_title(title)

        # --- 問題パート ---
        global_index = 0

        for ps in problem_sets:
            self._add_section_header(f"問題 {ps.original.number} の類題群", level=2)

            if include_originals:
                p = self.doc.add_paragraph()
                run = p.add_run("【元の問題】")
                run.bold = True
                run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
                p = self.doc.add_paragraph(ps.original.text)
                p_fmt = p.paragraph_format
                p_fmt.left_indent = Cm(0.5)
                self.doc.add_paragraph()  # 空行

            for gen in ps.variants:
                global_index += 1
                self._add_problem(gen, global_index)

            self.doc.add_paragraph()  # 空行

        # --- 改ページして解答・解説パート ---
        self.doc.add_page_break()
        self._add_title("解答・解説")

        global_index = 0
        for ps in problem_sets:
            self._add_section_header(
                f"問題 {ps.original.number} の解答・解説", level=2
            )

            for gen in ps.variants:
                global_index += 1
                self._add_answer_section(gen, global_index)

        # 保存
        output_path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(output_path))
        return output_path
